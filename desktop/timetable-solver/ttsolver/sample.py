"""Writes a timetable-constraints.xlsx in the portal's format, and makes realistic test instances."""
from __future__ import annotations

import random

from openpyxl import Workbook

from .model import Problem, Room, Slot, Unavailability, to_hours


def write_constraints(path: str, problem: Problem, extra_column: bool = False) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Sections"
    head = ["SlotIndex", "ScheduleSectionId", "BatchId", "InstructorId", "RoomTypeNeeded", "DurationHours", "StudentCount", "AllowedDays",
            "DayStartHour", "DayEndHour", "Label", "InstructorName", "BatchLabels", "PreferredDays"]
    ws.append(head + (["SomeNewerColumn"] if extra_column else []))
    for s in problem.slots:
        ws.append([s.index, s.section_id, ",".join(s.batch_ids), s.instructor_id, s.room_type, to_hours(s.duration), s.students, ",".join(s.allowed_days),
                   to_hours(s.day_start), to_hours(s.day_end), s.label, s.instructor_name, "; ".join(s.batch_labels), ",".join(s.preferred_days)]
                  + (["ignored"] if extra_column else []))
    r = wb.create_sheet("Rooms")
    r.append(["RoomId", "Type", "Capacity", "RoomName"])
    for k in problem.rooms:
        r.append([k.id, k.type, k.capacity, k.name])
    u = wb.create_sheet("Unavailability")
    u.append(["FacultyId", "DayOfWeek", "StartHour", "EndHour"])
    for x in problem.unavailability:
        u.append([x.faculty_id, x.day, to_hours(x.start), to_hours(x.end)])
    wb.save(path)


def generated_problem(n_slots: int = 600, n_rooms: int = 40, n_teachers: int = 80, n_batches: int = 36, seed: int = 7) -> Problem:
    """A university-like week: Mon–Fri 08:00–16:00, theory sections of 2–3 sessions (1 h or 1.5 h), 3-hour labs,
    some combined classes across two batches, teachers with a few hours away and a Friday prayer break."""
    rnd = random.Random(seed)
    days = ["Mon", "Tue", "Wed", "Thu", "Fri"]
    n_labs = max(4, n_rooms // 5)
    rooms = [Room(f"R{i:02d}", "THEORY", rnd.choice([40, 50, 60, 60, 80, 120]), f"Room {i:02d}") for i in range(n_rooms - n_labs)]
    rooms += [Room(f"L{i:02d}", "LAB", rnd.choice([30, 40, 40, 50]), f"Lab {i:02d}") for i in range(n_labs)]
    teachers = [f"T{i:02d}" for i in range(n_teachers)]
    batches = [(f"B{i:02d}", rnd.choice([30, 35, 40, 45, 50])) for i in range(n_batches)]
    slots: list[Slot] = []
    load = {t: 0 for t in teachers}
    sec = 0
    while len(slots) < n_slots:
        b, size = batches[sec % n_batches]
        combined = rnd.random() < 0.08
        bids = [b] + ([batches[(sec + 1) % n_batches][0]] if combined else [])
        students = size + (batches[(sec + 1) % n_batches][1] if combined else 0)
        lab = rnd.random() < 0.15 and not combined
        teacher = min(teachers, key=lambda t: (load[t], rnd.random()))
        if lab:
            n, dur, rtype, students = 1, 6, "LAB", min(students, 40)
        else:
            (n, dur), rtype = rnd.choice([(2, 3), (3, 2), (2, 2)]), "THEORY"  # sessions a week, half-hours each
        for _ in range(n):
            if len(slots) >= n_slots:
                break
            slots.append(Slot(len(slots), f"S{sec}", bids, teacher, rtype, dur, students, days, 16, 32,
                              label=f"C{sec:03d} {'Lab' if lab else 'A'}", instructor_name=f"Teacher {teacher}", batch_labels=[f"Batch {x}" for x in bids],
                              preferred_days=rnd.sample(days, 3) if rnd.random() < 0.3 else []))
            load[teacher] += dur
        sec += 1
    unav = [Unavailability(t, "Fri", 26, 28) for t in teachers]  # Friday prayer break 13:00–14:00
    for t in rnd.sample(teachers, n_teachers // 4):
        unav.append(Unavailability(t, rnd.choice(days), 16, 20))
    return Problem(slots, rooms, unav)


def tiny_problem() -> Problem:
    """Two batches, a combined class, a lab, one busy teacher — with a known clash-free answer."""
    D = ["Mon", "Tue", "Wed"]
    rooms = [Room("R1", "THEORY", 60, "Room 1"), Room("R2", "THEORY", 40, "Room 2"), Room("L1", "LAB", 30, "Lab 1")]

    def S(i, sec, b, t, rt, dur, n, label):
        return Slot(i, sec, b, t, rt, dur, n, D, 16, 26, label=label, instructor_name=t)

    slots = [
        S(0, "CS101-A", ["B1"], "Khan", "THEORY", 2, 35, label="CS101 A"), S(1, "CS101-A", ["B1"], "Khan", "THEORY", 2, 35, label="CS101 A"),
        S(2, "MA101-A", ["B1", "B2"], "Ali", "THEORY", 3, 55, label="MA101 A (combined)"), S(3, "MA101-A", ["B1", "B2"], "Ali", "THEORY", 3, 55, label="MA101 A (combined)"),
        S(4, "CS102-L", ["B2"], "Khan", "LAB", 6, 25, label="CS102 Lab"),
        S(5, "EN101-A", ["B2"], "Sara", "THEORY", 2, 30, label="EN101 A"),
    ]
    return Problem(slots, rooms, [Unavailability("Khan", "Mon", 16, 20)])
