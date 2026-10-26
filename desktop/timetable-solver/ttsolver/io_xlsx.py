"""Reading timetable-constraints.xlsx and writing / reading timetable-solution.xlsx.

The constraints file is read by sheet and column NAME (never position), unknown extra columns are ignored, and every
problem is listed with its sheet and row. The solution file follows exactly what the portal's upload route reads:
sheet "Solution" (row 1 header; columns 1-4 = SlotIndex, Day, StartHour, RoomId) and sheet "Summary" (B1 hard
violations, B2 solver branches, B3 objective score).
"""
from __future__ import annotations

import math
from dataclasses import dataclass, field
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .model import DAY_ORDER, Assignment, Placement, Problem, Room, Slot, Unavailability, fmt_time, to_half, to_hours

REQUIRED = {
    "Sections": ["SlotIndex", "ScheduleSectionId", "BatchId", "InstructorId", "RoomTypeNeeded", "DurationHours", "StudentCount", "AllowedDays", "DayStartHour", "DayEndHour"],
    "Rooms": ["RoomId", "Type", "Capacity"],
    "Unavailability": ["FacultyId", "DayOfWeek", "StartHour", "EndHour"],
}
OPTIONAL = {"Sections": ["Label", "InstructorName", "BatchLabels", "PreferredDays"], "Rooms": ["RoomName"]}


@dataclass
class LoadResult:
    problem: Problem | None
    errors: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)


def _norm(h: Any) -> str:
    return str(h or "").strip().lower().replace(" ", "")


def _cell(v: Any) -> Any:
    if isinstance(v, str):
        v = v.strip()
        return v if v else None
    return v


def _rows(ws, sheet: str, errors: list[str]):
    """Yields (row number, {column name: value}) using the header row; reports missing required columns."""
    it = ws.iter_rows(values_only=True)
    try:
        header = next(it)
    except StopIteration:
        errors.append(f"Sheet {sheet}: it is empty (no header row).")
        return
    pos = {}
    for i, h in enumerate(header or []):
        n = _norm(h)
        if n and n not in pos:
            pos[n] = i
    missing = [c for c in REQUIRED[sheet] if _norm(c) not in pos]
    if missing:
        errors.append(f"Sheet {sheet}: missing column(s) {', '.join(missing)}.")
        return
    wanted = REQUIRED[sheet] + OPTIONAL.get(sheet, [])
    for r, values in enumerate(it, start=2):
        if values is None or all(_cell(v) is None for v in values):
            continue
        row = {}
        for c in wanted:
            i = pos.get(_norm(c))
            row[c] = _cell(values[i]) if i is not None and i < len(values) else None
        yield r, row


def _num(row: dict, col: str, where: str, errors: list[str], *, integer=False, minimum: float | None = None) -> float | None:
    v = row.get(col)
    if v is None:
        errors.append(f"{where}: {col} empty")
        return None
    try:
        x = float(v)
    except (TypeError, ValueError):
        errors.append(f"{where}: {col} is not a number ({v!r})")
        return None
    if integer and x != int(x):
        errors.append(f"{where}: {col} must be a whole number ({v!r})")
        return None
    if minimum is not None and x < minimum:
        errors.append(f"{where}: {col} must be at least {minimum:g} ({v!r})")
        return None
    return x


def _on_grid(x: float) -> bool:
    return abs(x * 2 - round(x * 2)) < 1e-6


def _days(text: Any, where: str, col: str, errors: list[str], required=True) -> list[str]:
    if text is None:
        if required:
            errors.append(f"{where}: {col} empty")
        return []
    out = []
    for d in str(text).replace(";", ",").split(","):
        d = d.strip()
        if not d:
            continue
        canon = next((x for x in DAY_ORDER if x.lower() == d[:3].lower()), None)
        if not canon:
            errors.append(f"{where}: {col} has an unknown day {d!r} (use Mon,Tue,Wed,Thu,Fri,Sat,Sun)")
            continue
        if canon not in out:
            out.append(canon)
    if required and not out:
        errors.append(f"{where}: {col} lists no days")
    return out


def read_constraints(path: str) -> LoadResult:
    errors: list[str] = []
    warnings: list[str] = []
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
    except Exception as e:  # noqa: BLE001 — any unreadable file
        return LoadResult(None, [f"Can't open the file: {e}"])
    sheets = {_norm(n): n for n in wb.sheetnames}
    for s in REQUIRED:
        if _norm(s) not in sheets:
            errors.append(f"The file has no sheet named {s}.")
    if errors:
        return LoadResult(None, errors)

    slots: list[Slot] = []
    seen_index: set[int] = set()
    for r, row in _rows(wb[sheets["sections"]], "Sections", errors):
        where = f"Sections row {r}"
        n0 = len(errors)
        idx = _num(row, "SlotIndex", where, errors, integer=True, minimum=0)
        dur = _num(row, "DurationHours", where, errors, minimum=0.5)
        cnt = _num(row, "StudentCount", where, errors, integer=True, minimum=0)
        ds = _num(row, "DayStartHour", where, errors, minimum=0)
        de = _num(row, "DayEndHour", where, errors, minimum=0)
        for col in ("ScheduleSectionId", "BatchId", "InstructorId", "RoomTypeNeeded"):
            if row.get(col) is None:
                errors.append(f"{where}: {col} empty")
        days = _days(row.get("AllowedDays"), where, "AllowedDays", errors)
        pref = _days(row.get("PreferredDays"), where, "PreferredDays", errors, required=False)
        if dur is not None and not _on_grid(dur):
            errors.append(f"{where}: DurationHours {dur:g} is not in half-hour steps")
        for col, x in (("DayStartHour", ds), ("DayEndHour", de)):
            if x is not None and not _on_grid(x):
                errors.append(f"{where}: {col} {x:g} is not on the half-hour grid")
        if ds is not None and de is not None and de <= ds:
            errors.append(f"{where}: DayEndHour ({de:g}) is not after DayStartHour ({ds:g})")
        if idx is not None and int(idx) in seen_index:
            errors.append(f"{where}: SlotIndex {int(idx)} appears twice")
        if len(errors) > n0:
            continue
        seen_index.add(int(idx))
        batch_ids = [b.strip() for b in str(row["BatchId"]).split(",") if b.strip()]
        labels = [x.strip() for x in str(row.get("BatchLabels") or "").split(";")]
        slots.append(Slot(
            index=int(idx), section_id=str(row["ScheduleSectionId"]), batch_ids=batch_ids, instructor_id=str(row["InstructorId"]),
            room_type=str(row["RoomTypeNeeded"]), duration=to_half(dur), students=int(cnt), allowed_days=days,
            day_start=to_half(ds), day_end=to_half(de), label=str(row.get("Label") or ""), instructor_name=str(row.get("InstructorName") or ""),
            batch_labels=labels if len(labels) == len(batch_ids) else [], preferred_days=pref, row=r,
        ))

    rooms: list[Room] = []
    seen_room: set[str] = set()
    for r, row in _rows(wb[sheets["rooms"]], "Rooms", errors):
        where = f"Rooms row {r}"
        n0 = len(errors)
        cap = _num(row, "Capacity", where, errors, integer=True, minimum=0)
        for col in ("RoomId", "Type"):
            if row.get(col) is None:
                errors.append(f"{where}: {col} empty")
        if row.get("RoomId") is not None and str(row["RoomId"]) in seen_room:
            errors.append(f"{where}: RoomId {row['RoomId']} appears twice")
        if len(errors) > n0:
            continue
        seen_room.add(str(row["RoomId"]))
        rooms.append(Room(id=str(row["RoomId"]), type=str(row["Type"]), capacity=int(cap), name=str(row.get("RoomName") or "")))

    unav: list[Unavailability] = []
    for r, row in _rows(wb[sheets["unavailability"]], "Unavailability", errors):
        where = f"Unavailability row {r}"
        n0 = len(errors)
        st = _num(row, "StartHour", where, errors, minimum=0)
        en = _num(row, "EndHour", where, errors, minimum=0)
        if row.get("FacultyId") is None:
            errors.append(f"{where}: FacultyId empty")
        days = _days(row.get("DayOfWeek"), where, "DayOfWeek", errors)
        if st is not None and en is not None and en <= st:
            errors.append(f"{where}: EndHour ({en:g}) is not after StartHour ({st:g})")
        if len(errors) > n0:
            continue
        for d in days:
            # Off-grid times are widened to whole half-hours, so the teacher stays free for the whole window.
            unav.append(Unavailability(str(row["FacultyId"]), d, math.floor(float(st) * 2 + 1e-9), math.ceil(float(en) * 2 - 1e-9)))

    wb.close()
    if not slots and not errors:
        errors.append("Sheet Sections has no sessions to schedule.")
    if not rooms and not errors:
        errors.append("Sheet Rooms has no rooms.")
    if errors:
        return LoadResult(None, errors, warnings)
    room_types = {r.type for r in rooms}
    for s in slots:
        if s.room_type not in room_types:
            warnings.append(f"Sections row {s.row}: no room of type {s.room_type!r} exists ({s.name})")
    return LoadResult(Problem(slots, rooms, unav, source=path), errors, warnings)


# ------------------------------------------------------------------------------------------------ solution file
HEADER_FILL = PatternFill("solid", fgColor="0D2A4A")
HEADER_FONT = Font(bold=True, color="FFFFFF")


def write_solution(path: str, problem: Problem, assignment: Assignment, *, hard_violations: int, branches: int, score: float, grids: bool = True) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Solution"
    ws.append(["SlotIndex", "Day", "StartHour", "RoomId", "Label", "Instructor", "Batches", "Room name", "End"])
    for c in ws[1]:
        c.fill, c.font = HEADER_FILL, HEADER_FONT
    for s in sorted(problem.slots, key=lambda x: x.index):
        p = assignment[s.index]
        room = problem.room_by_id.get(p.room_id)
        ws.append([s.index, p.day, to_hours(p.start), p.room_id, s.label, s.teacher, "; ".join(s.batch_name(b) for b in s.batch_ids), room.label if room else "", fmt_time(p.start + s.duration)])
    for col, w in zip("ABCDEFGHI", (10, 7, 10, 28, 30, 24, 40, 18, 8)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"

    sm = wb.create_sheet("Summary")
    sm["A1"], sm["B1"] = "Hard violations", int(hard_violations)
    sm["A2"], sm["B2"] = "Solver branches", int(branches)
    sm["A3"], sm["B3"] = "Objective score (lower is better)", float(score)
    sm["A4"], sm["B4"] = "Slots", len(problem.slots)
    sm.column_dimensions["A"].width = 34
    for c in ("A1", "A2", "A3", "A4"):
        sm[c].font = Font(bold=True)

    if grids:
        from .report import grid_blocks
        for title, kind in (("By batch", "batch"), ("By instructor", "instructor"), ("By room", "room")):
            _grid_sheet(wb.create_sheet(title), grid_blocks(problem, assignment, kind))
    wb.save(path)


def _grid_sheet(ws, blocks) -> None:
    thin = Alignment(wrap_text=True, vertical="top")
    r = 1
    for b in blocks:
        ws.cell(r, 1, b.title).font = Font(bold=True, size=12)
        r += 1
        ws.cell(r, 1, "Time").font = Font(bold=True)
        for j, d in enumerate(b.days, start=2):
            ws.cell(r, j, d).font = Font(bold=True)
        r += 1
        top = r
        for i, t in enumerate(b.times):
            ws.cell(top + i, 1, fmt_time(t))
        for e in b.entries:
            col = 2 + b.days.index(e.day)
            row = top + b.times.index(e.start)
            cell = ws.cell(row, col, e.text)
            cell.alignment = thin
            cell.fill = PatternFill("solid", fgColor=e.color.lstrip("#"))
            span = max(1, e.length)
            if span > 1:
                ws.merge_cells(start_row=row, start_column=col, end_row=row + span - 1, end_column=col)
        r = top + len(b.times) + 2
    ws.column_dimensions["A"].width = 8
    for j in range(2, 9):
        ws.column_dimensions[chr(64 + j)].width = 26


def read_solution(path: str) -> tuple[dict[int, dict[str, Any]], dict[str, float]]:
    """Reads a solution the way the portal's upload route does: rows 2.. of sheet Solution, columns 1-4, a row counts only
    when SlotIndex and StartHour are numbers and Day and RoomId are not empty; Summary B1-B3."""
    wb = load_workbook(path, data_only=True)
    if "Solution" not in wb.sheetnames:
        raise ValueError('this file has no "Solution" sheet')
    ws = wb["Solution"]
    out: dict[int, dict[str, Any]] = {}
    for r, row in enumerate(ws.iter_rows(values_only=True), start=1):
        if r == 1 or row is None:
            continue
        vals = list(row) + [None] * 4
        try:
            idx = float(vals[0])
            start = float(vals[2])
        except (TypeError, ValueError):
            continue
        day, room = str(vals[1] or ""), str(vals[3] or "")
        if idx != idx or start != start or not day or not room:
            continue
        out[int(idx)] = {"day": day, "startHour": start, "roomId": room}
    summary = {"hardViolations": 0.0, "generations": 0.0, "score": 0.0}
    if "Summary" in wb.sheetnames:
        sm = wb["Summary"]
        for key, ref in (("hardViolations", "B1"), ("generations", "B2"), ("score", "B3")):
            try:
                summary[key] = float(sm[ref].value or 0)
            except (TypeError, ValueError):
                summary[key] = 0.0
    return out, summary


def assignment_from_solution(rows: dict[int, dict[str, Any]]) -> Assignment:
    out: Assignment = {}
    for i, v in rows.items():
        raw = float(v["startHour"]) * 2
        # An off-grid start (e.g. 8.25) is kept as it is, so the checker reports it instead of silently rounding it.
        out[i] = Placement(v["day"], int(round(raw)) if abs(raw - round(raw)) < 1e-9 else raw, v["roomId"])  # type: ignore[arg-type]
    return out
