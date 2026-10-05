"""
Checks a finished solution file against the constraints file it came from,
using separate, from-scratch code (not the solver's own checker), and
replays what the portal's "Upload Solution" step reads.

    python test_roundtrip.py test-constraints.xlsx timetable-solution.xlsx
"""
import sys
from collections import defaultdict
import openpyxl


def main(constraints, solution):
    cw = openpyxl.load_workbook(constraints, data_only=True)
    sw = openpyxl.load_workbook(solution, data_only=True)
    problems = []

    # --- exactly what upload-solution/route.ts reads ---
    if "Solution" not in sw.sheetnames:
        print('FAIL: no "Solution" sheet'); return 1
    sol = {}
    for rn, row in enumerate(sw["Solution"].iter_rows(values_only=True), start=1):
        if rn == 1:
            continue
        idx, day, start, rid = row[0], row[1], row[2], row[3]
        if idx is None or not day or start is None or not rid:
            continue
        sol[int(idx)] = (str(day), float(start), str(rid))
    sm = sw["Summary"]
    b1, b2, b3 = sm["B1"].value, sm["B2"].value, sm["B3"].value
    print(f"portal would read: {len(sol)} entries, hardViolations={b1}, generations={b2}, score={b3}")

    # --- constraints, re-read independently ---
    secs = list(cw["Sections"].iter_rows(min_row=2, values_only=True))
    rooms = {r[0]: (r[1], int(r[2])) for r in cw["Rooms"].iter_rows(min_row=2, values_only=True) if r[0]}
    un = [(u[0], u[1], float(u[2]), float(u[3])) for u in cw["Unavailability"].iter_rows(min_row=2, values_only=True) if u[0]]
    if len(sol) != len(secs):
        problems.append(f"entry count {len(sol)} != slots {len(secs)} (portal would reject this)")

    cnt = defaultdict(int)
    by_day = defaultdict(list)
    for r in secs:
        idx, instr, rtype, dur, stu = int(r[0]), r[3], r[4], float(r[5]), int(r[6])
        days, ds, de = str(r[7]).split(","), float(r[8]), float(r[9])
        batches = set(str(r[2]).split(","))
        day, start, rid = sol[idx]
        end = start + dur
        if day not in days: cnt["day not allowed"] += 1
        if start < ds or end > de: cnt["outside hours"] += 1
        if rid not in rooms: cnt["unknown room"] += 1; continue
        if rooms[rid][1] < stu: cnt["room too small"] += 1
        if rooms[rid][0] != rtype: cnt[f"{rtype} class in {rooms[rid][0]} room"] += 1
        for f, d, a, b in un:
            if f == instr and d == day and start < b and a < end: cnt["instructor blocked time"] += 1
        by_day[day].append((idx, start, end, instr, rid, batches))
    for day, items in by_day.items():
        for i in range(len(items)):
            for j in range(i + 1, len(items)):
                a, b = items[i], items[j]
                if not (a[1] < b[2] and b[1] < a[2]): continue
                if a[3] == b[3]: cnt["instructor double-booked"] += 1
                if a[4] == b[4]: cnt["room double-booked"] += 1
                if a[5] & b[5]: cnt["batch double-booked"] += 1
    print("independent check:", dict(cnt) if cnt else "NO VIOLATIONS")
    for p in problems: print("PROBLEM:", p)
    hard = sum(v for k, v in cnt.items() if "class in" not in k)
    print("matches Summary!B1:", hard == b1 or f"(file says {b1}, checker counted {hard} excluding intentional theory-in-lab)")
    return 0 if not problems and hard == 0 else 2


if __name__ == "__main__":
    sys.exit(main(sys.argv[1], sys.argv[2]))
