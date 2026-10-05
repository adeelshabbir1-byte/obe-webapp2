"""
Makes a made-up timetable-constraints.xlsx (same layout the portal exports)
so you can try the solver without touching real data.

    python make_test_data.py            # -> test-constraints.xlsx, shaped like a
                                        #    mid-size department: ~182 theory sections,
                                        #    25 labs, 10 lecture rooms, 5 labs
    python make_test_data.py --small    # a quick, easy one
"""
import argparse
import json
import random

import openpyxl

DAYS6 = "Mon,Tue,Wed,Thu,Fri,Sat"


def build(seed=11, small=False, lecture_rooms=None):
    rnd = random.Random(seed)
    n_batches = 6 if small else 24
    n_theory = 30 if small else 182
    n_lab = 4 if small else 25
    n_instr = 14 if small else 60
    n_lec, n_labrooms = (3, 2) if small else (10, 5)
    if lecture_rooms:
        n_lec = lecture_rooms
    n_clubbed = 3 if small else 15

    programs = ["AI", "SE", "CY", "CS"]
    batches = [{"id": f"b{i}", "label": f"BS {programs[i % 4]} - Term {i // 4 + 1}", "students": rnd.randint(38, 56)} for i in range(n_batches)]
    instr = [{"id": f"i{i}", "name": f"Instructor {i + 1}"} for i in range(n_instr)]
    rooms = [{"id": f"L{i}", "name": f"Room {101 + i}", "type": "LECTURE", "cap": rnd.choice([60, 65, 70])} for i in range(n_lec)]
    rooms += [{"id": f"X{i}", "name": f"Lab {i + 1}", "type": "LAB", "cap": 60} for i in range(n_labrooms)]

    slots = []
    sec = 0

    def add_section(kind, covered, label):
        nonlocal sec
        sec += 1
        ins = rnd.choice(instr)
        sessions, dur = (2, 1.5) if kind == "LECTURE" else (1, 3.0)
        total = sum(b["students"] for b in covered)
        for _ in range(sessions):
            slots.append({
                "section": f"s{sec}", "batches": [b["id"] for b in covered], "instr": ins["id"], "instr_name": ins["name"],
                "type": kind, "dur": dur, "students": total if len(covered) == 1 else (total + 1) // 2,
                "label": label, "batch_label": "; ".join(b["label"] for b in covered),
            })

    for k in range(n_clubbed):
        a, b = rnd.sample(batches, 2)
        add_section("LECTURE", [a, b], f"CLUB-{k + 1:02d} Section A (combined)")
    for k in range(n_theory - n_clubbed):
        b = batches[k % n_batches]
        add_section("LECTURE", [b], f"CS-{100 + k} Section A")
    for k in range(n_lab):
        b = batches[(k * 3) % n_batches]
        add_section("LAB", [b], f"CS-{100 + k}-L Section A")
    for i, s in enumerate(slots):
        s["index"] = i

    unavail = []
    for ins in rnd.sample(instr, n_instr // 2):
        d = rnd.choice(DAYS6.split(","))
        st = rnd.choice([8, 9, 10, 12, 13])
        unavail.append({"faculty": ins["id"], "day": d, "start": st, "end": st + rnd.choice([2, 3, 4])})
    return slots, rooms, unavail


def write_xlsx(path, slots, rooms, unavail):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sections"
    ws.append(["SlotIndex", "ScheduleSectionId", "BatchId", "InstructorId", "RoomTypeNeeded", "DurationHours", "StudentCount",
               "AllowedDays", "DayStartHour", "DayEndHour", "Label", "InstructorName", "BatchLabels"])
    for s in slots:
        ws.append([s["index"], s["section"], ",".join(s["batches"]), s["instr"], s["type"], s["dur"], s["students"],
                   DAYS6, 8, 16, s["label"], s["instr_name"], s["batch_label"]])
    ws = wb.create_sheet("Rooms")
    ws.append(["RoomId", "Type", "Capacity", "RoomName"])
    for r in rooms:
        ws.append([r["id"], r["type"], r["cap"], r["name"]])
    ws = wb.create_sheet("Unavailability")
    ws.append(["FacultyId", "DayOfWeek", "StartHour", "EndHour"])
    for u in unavail:
        ws.append([u["faculty"], u["day"], u["start"], u["end"]])
    wb.save(path)


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--small", action="store_true")
    ap.add_argument("--lecture-rooms", type=int, help="override the number of lecture rooms")
    ap.add_argument("-o", "--output", default="test-constraints.xlsx")
    ap.add_argument("--json", help="also dump the data as JSON (used to benchmark the old web generator)")
    a = ap.parse_args()
    slots, rooms, unavail = build(small=a.small, lecture_rooms=a.lecture_rooms)
    write_xlsx(a.output, slots, rooms, unavail)
    if a.json:
        json.dump({"slots": slots, "rooms": rooms, "unavail": unavail}, open(a.json, "w"))
    print(f"wrote {a.output}: {len(slots)} class sessions, {len(rooms)} rooms, {len(unavail)} blocked instructor times")
